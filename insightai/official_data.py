from __future__ import annotations

import hashlib
import io
import json
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import Callable

import duckdb
import pandas as pd
from openpyxl import load_workbook


STRUCTURED = {
    "medical": {"label": "医疗", "workbook": "Healthcare_Analytics_Competition.xlsx", "tasks": "merged_problems2_task.json"},
    "telecom": {"label": "通信", "workbook": "telecom_operations_db.xlsx", "tasks": "merged_problems3_task.json"},
    "finance": {"label": "金融", "workbook": "financial_asset_management.xlsx", "tasks": "merged_problems1_task.json"},
}
UNSTRUCTURED = {
    "multi_step": {"label": "多步检索", "stem": "Multi-step_Retrieval"},
    "domain_calc": {"label": "领域计算", "stem": "Table_Domain-specific_Operations"},
    "table_query": {"label": "表格查询", "stem": "Table_Query"},
}
EXPECTED_TASKS = {"medical": 210, "telecom": 191, "finance": 279, "multi_step": 48, "domain_calc": 239, "table_query": 500}
SENSITIVE_HINTS = ("id_card", "身份证", "phone", "mobile", "电话", "手机", "email", "邮箱", "address", "地址", "住址", "contact")


def fix_zip_name(name: str) -> str:
    """Repair archives whose UTF-8 names were decoded as CP437."""
    try:
        return name.encode("cp437").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return name


def _find_inner_zip(source: Path, work: Path) -> Path:
    with zipfile.ZipFile(source) as outer:
        candidates = [item for item in outer.infolist() if not item.is_dir() and item.filename.lower().endswith(".zip")]
        if not candidates:
            raise ValueError("外层压缩包中未找到赛题 ZIP")
        # Task 1 is the largest nested archive in the organizer package.
        item = max(candidates, key=lambda value: value.file_size)
        target = work / "task1_official.zip"
        with outer.open(item) as src, target.open("wb") as dst:
            shutil.copyfileobj(src, dst, 8 * 1024 * 1024)
        return target


def resolve_task1_zip(source: str | Path, work: Path) -> Path:
    source = Path(source).resolve()
    if not zipfile.is_zipfile(source):
        raise ValueError(f"不是有效 ZIP：{source}")
    with zipfile.ZipFile(source) as archive:
        names = [fix_zip_name(item.filename) for item in archive.infolist()]
        if any("merged_problems" in name for name in names):
            return source
    return _find_inner_zip(source, work)


def _member(archive: zipfile.ZipFile, *, suffix: str, contains: str | None = None) -> zipfile.ZipInfo:
    found = []
    for item in archive.infolist():
        name = fix_zip_name(item.filename)
        if name.endswith(suffix) and (contains is None or contains in name):
            found.append(item)
    if len(found) != 1:
        raise ValueError(f"ZIP 内匹配失败：suffix={suffix}, contains={contains}, found={len(found)}")
    return found[0]


def _extract_member(archive: zipfile.ZipFile, item: zipfile.ZipInfo, target: Path) -> Path:
    target.parent.mkdir(parents=True, exist_ok=True)
    with archive.open(item) as src, target.open("wb") as dst:
        shutil.copyfileobj(src, dst, 8 * 1024 * 1024)
    return target


def _clean_value(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value


def _stream_sheet(connection: duckdb.DuckDBPyConnection, workbook: Path, sheet: str, batch_size: int = 10_000) -> tuple[int, list[str]]:
    book = load_workbook(workbook, read_only=True, data_only=True)
    page = book[sheet]
    rows = page.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else f"column_{index + 1}" for index, value in enumerate(next(rows))]
    total, created, batch = 0, False, []
    quoted = '"' + sheet.replace('"', '""') + '"'

    def flush(records: list[tuple]) -> int:
        nonlocal created
        frame = pd.DataFrame.from_records(records, columns=headers)
        connection.register("_insightai_batch", frame)
        if not created:
            connection.execute(f"CREATE OR REPLACE TABLE {quoted} AS SELECT * FROM _insightai_batch")
            created = True
        else:
            connection.execute(f"INSERT INTO {quoted} BY NAME SELECT * FROM _insightai_batch")
        connection.unregister("_insightai_batch")
        return len(records)

    for row in rows:
        batch.append(tuple(_clean_value(value) for value in row))
        if len(batch) >= batch_size:
            total += flush(batch)
            batch = []
    if batch or not created:
        total += flush(batch)
    book.close()
    return total, headers


TASK_BLOCK_RE = re.compile(r"(?ms)^===\s*([^=\r\n]+?)\s*===\s*\r?\n(.*?)(?=^===\s*[^=\r\n]+?\s*===\s*$|\Z)")


def index_markdown(text: str) -> dict[str, str]:
    return {match.group(1).strip(): match.group(2).strip() for match in TASK_BLOCK_RE.finditer(text)}


def parse_table_block(block: str) -> list[pd.DataFrame]:
    cleaned = re.sub(r"^```(?:markdown|html)?\s*|\s*```$", "", block.strip(), flags=re.I | re.S)
    frames: list[pd.DataFrame] = []
    if "<table" in cleaned.lower():
        try:
            frames.extend(pd.read_html(io.StringIO(cleaned)))
        except ValueError:
            pass
    lines = [line.strip() for line in cleaned.splitlines() if line.strip().startswith("|")]
    groups: list[list[str]] = []
    current: list[str] = []
    def split_row(line: str) -> list[str]:
        body = line[1:] if line.startswith("|") else line
        body = body[:-1] if body.endswith("|") else body
        return [cell.strip() for cell in body.split("|")]

    for line in lines:
        if current and len(split_row(line)) != len(split_row(current[0])):
            groups.append(current)
            current = []
        current.append(line)
    if current:
        groups.append(current)
    for group in groups:
        rows = [split_row(line) for line in group]
        separator = len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in rows[1])
        if separator:
            frames.append(pd.DataFrame(rows[2:], columns=rows[0]))
    return frames


def _read_json(archive: zipfile.ZipFile, item: zipfile.ZipInfo):
    return json.loads(archive.read(item).decode("utf-8-sig"))


def _write_json(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def load_data_dictionary(path: str | Path) -> dict:
    """Normalize the organizer's three data-dictionary sheets by column position."""
    workbook = pd.ExcelFile(path)
    result = {}
    for sheet in workbook.sheet_names:
        frame = pd.read_excel(path, sheet_name=sheet)
        records = []
        for row in frame.itertuples(index=False, name=None):
            values = [(None if pd.isna(value) else str(value).strip()) for value in row]
            if len(values) < 8 or not values[2] or not values[5]:
                continue
            records.append({
                "database_cn": values[0], "database_en": values[1],
                "table_en": values[2], "table_cn": values[3], "table_comment": values[4],
                "field_en": values[5], "field_cn": values[6], "field_comment": values[7],
            })
        label = str(sheet)
        key = next((name for name, config in STRUCTURED.items() if config["label"] == label), label)
        result[key] = records
    return result


def infer_relationships(tables: list[dict]) -> list[dict]:
    by_field: dict[str, list[str]] = {}
    for table in tables:
        for field in table["columns"]:
            if str(field).lower().endswith("_id"):
                by_field.setdefault(str(field), []).append(table["name"])
    return [{"field": field, "tables": names} for field, names in sorted(by_field.items()) if len(names) > 1]


def prepare_official_data(source: str | Path, output: str | Path, progress: Callable[[str], None] | None = None) -> dict:
    output = Path(output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    raw = output / "raw"
    raw.mkdir(exist_ok=True)
    notify = progress or (lambda message: None)
    with tempfile.TemporaryDirectory(prefix="insightai_zip_") as temp_name:
        task_zip = resolve_task1_zip(source, Path(temp_name))
        packaged = raw / "赛题1_智能数据分析系统.zip"
        if Path(task_zip).resolve() != packaged.resolve():
            shutil.copy2(task_zip, packaged)
        manifest = {"source_sha256": hashlib.sha256(packaged.read_bytes()).hexdigest(), "structured": {}, "unstructured": {}, "expected_tasks": EXPECTED_TASKS}
        with zipfile.ZipFile(task_zip) as archive:
            dictionary_path = _extract_member(archive, _member(archive, suffix="数据表结构.xlsx"), raw / "数据表结构.xlsx")
            dictionary = load_data_dictionary(dictionary_path)
            _write_json(output / "schema_dictionary.json", dictionary)
            for key, config in STRUCTURED.items():
                notify(f"准备{config['label']}数据")
                workbook = _extract_member(archive, _member(archive, suffix=config["workbook"]), raw / key / config["workbook"])
                tasks = _read_json(archive, _member(archive, suffix=config["tasks"]))
                samples = _read_json(archive, _member(archive, suffix="样例答案.json", contains=f"/{config['label']}/"))
                _write_json(output / "tasks" / f"{key}.json", tasks)
                _write_json(output / "samples" / f"{key}.json", samples)
                database = output / "databases" / f"{key}.duckdb"
                database.parent.mkdir(exist_ok=True)
                database.unlink(missing_ok=True)
                connection = duckdb.connect(str(database))
                book = load_workbook(workbook, read_only=True, data_only=True)
                sheet_names = list(book.sheetnames)
                book.close()
                tables = []
                for sheet in sheet_names:
                    notify(f"{config['label']}：导入 {sheet}")
                    rows, columns = _stream_sheet(connection, workbook, sheet)
                    tables.append({"name": sheet, "rows": rows, "columns": columns})
                connection.close()
                manifest["structured"][key] = {"label": config["label"], "tasks": len(tasks), "database": str(database.relative_to(output)), "tables": tables, "rows": sum(table["rows"] for table in tables), "dictionary_fields": len(dictionary.get(key, [])), "relationships": infer_relationships(tables)}
            for key, config in UNSTRUCTURED.items():
                notify(f"索引{config['label']}任务")
                tasks = _read_json(archive, _member(archive, suffix=f"{config['stem']}_task.json"))
                text = archive.read(_member(archive, suffix=f"{config['stem']}-数据.md")).decode("utf-8-sig")
                blocks = index_markdown(text)
                _write_json(output / "tasks" / f"{key}.json", tasks)
                index_path = output / "indexes" / f"{key}.jsonl"
                index_path.parent.mkdir(exist_ok=True)
                with index_path.open("w", encoding="utf-8") as handle:
                    for task in tasks:
                        handle.write(json.dumps({"id": task["id"], "block": blocks.get(str(task["id"]), "")}, ensure_ascii=False) + "\n")
                manifest["unstructured"][key] = {"label": config["label"], "tasks": len(tasks), "blocks": sum(bool(blocks.get(str(task["id"]))) for task in tasks), "index": str(index_path.relative_to(output))}
            samples = _read_json(archive, _member(archive, suffix="样例答案.json", contains="非结构化数据/"))
            _write_json(output / "samples" / "unstructured.json", samples)
        manifest["total_tasks"] = sum(item["tasks"] for item in manifest["structured"].values()) + sum(item["tasks"] for item in manifest["unstructured"].values())
        _write_json(output / "manifest.json", manifest)
        notify("官方数据准备完成")
        return manifest


def load_manifest(cache: str | Path) -> dict | None:
    path = Path(cache) / "manifest.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else None


def load_tasks(cache: str | Path, dataset: str) -> list[dict]:
    return json.loads((Path(cache) / "tasks" / f"{dataset}.json").read_text(encoding="utf-8"))


def load_block_index(cache: str | Path, dataset: str) -> dict[str, str]:
    path = Path(cache) / "indexes" / f"{dataset}.jsonl"
    result = {}
    if path.exists():
        for line in path.read_text(encoding="utf-8").splitlines():
            item = json.loads(line)
            result[str(item["id"])] = item["block"]
    return result


def mask_value(column: str, value):
    if value is None or not any(hint.lower() in column.lower() for hint in SENSITIVE_HINTS):
        return value
    text = str(value)
    if "@" in text:
        head, _, domain = text.partition("@")
        return (head[:1] + "***@" + domain) if head else "***@" + domain
    if len(text) <= 4:
        return "*" * len(text)
    return text[:3] + "*" * min(8, len(text) - 5) + text[-2:]


def mask_frame(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    for column in result.columns:
        if any(hint.lower() in str(column).lower() for hint in SENSITIVE_HINTS):
            result[column] = result[column].map(lambda value: mask_value(str(column), value))
    return result
